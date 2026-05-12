import streamlit as st
import json
import os
import pandas as pd

# st.set_page_config(layout="wide")

# detect mobile
is_mobile = st.context.headers.get("user-agent", "").lower()

mobile = any(x in is_mobile for x in ["iphone", "android", "mobile"])

st.title("Danh sách sản phẩm")

if st.button("📥 Xuất Excel"):

    if os.path.exists("data.json"):

        with open("data.json", "r", encoding="utf-8") as f:
            data = json.load(f)

        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.drawing.image import Image as XLImage

        from openpyxl.drawing.spreadsheet_drawing import (
            OneCellAnchor,
            AnchorMarker
        )

        from openpyxl.drawing.xdr import (
            XDRPositiveSize2D
        )

        from openpyxl.utils.units import (
            pixels_to_EMU
        )

        wb = Workbook()
        ws = wb.active

        ws.title = "Sản phẩm"

        # ================= HEADER =================
        headers = [
            "Hình ảnh",
            "Tên",
            "Chiều dài",
            "Chiều rộng"
        ]

        ws.append(headers)

        thin = Side(
            style="thin",
            color="000000"
        )

        # style header
        for cell in ws[1]:

            cell.font = Font(
                bold=True,
                size=12
            )

            cell.font = Font(
                bold=True,
                size=12,
                color="FFFFFF"  # màu chữ trắng
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            # màu nền
            cell.fill = PatternFill(
                fill_type="solid",
                start_color="4F81BD",
                end_color="4F81BD"
            )
        
        for row_cells in ws.iter_rows():

            for cell in row_cells:

                cell.border = Border(
                    left=thin,
                    right=thin,
                    top=thin,
                    bottom=thin
                )

        # ================= DATA =================
        row = 2

        for item in data:

            # tên
            ws.cell(
                row=row,
                column=2,
                value=item["ten"]
            )

            # chiều dài
            ws.cell(
                row=row,
                column=3,
                value=item["chieu_dai"]
            )

            # chiều rộng
            ws.cell(
                row=row,
                column=4,
                value=item["chieu_rong"]
            )

            # căn giữa dữ liệu
            for col in range(2, 5):

                ws.cell(
                    row=row,
                    column=col
                ).alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

            # ================= IMAGE =================
            if os.path.exists(item["images"]):

                img = XLImage(item["images"])

                # size ảnh
                img.width = 120
                img.height = 100

                # center ảnh
                marker = AnchorMarker(
                    col=0,
                    colOff=pixels_to_EMU(5),
                    row=row - 1,
                    rowOff=pixels_to_EMU(5)
                )

                img.anchor = OneCellAnchor(
                    _from=marker,
                    ext=XDRPositiveSize2D(
                        pixels_to_EMU(img.width),
                        pixels_to_EMU(img.height)
                    )
                )

                ws.add_image(img)

            # chiều cao row
            ws.row_dimensions[row].height = 80

            row += 1

        # ================= COLUMN SIZE =================
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 18

        excel_file = "products.xlsx"

        wb.save(excel_file)

        with open(excel_file, "rb") as file:

            st.download_button(
                label="📥 Tải Excel",
                data=file,
                file_name="products.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
if os.path.exists("data.json"):

    with open("data.json", "r", encoding="utf-8") as f:
        data = json.load(f)

    if len(data) > 0:

        # ================= MOBILE =================
        if mobile:

            # st.subheader("Mobile View")

            for i in range(0, len(data), 2):

                cols = st.columns(2)

                for j in range(2):

                    if i + j < len(data):

                        index = i + j
                        item = data[index]

                        with cols[j]:

                            with st.container(border=True):

                                if os.path.exists(item["images"]):
                                    st.image(
                                        item["images"],
                                        use_container_width=True
                                    )

                                st.write(f"### {item['ten']}")

                                col1, col2 = st.columns(2)

                                col1.metric(
                                    "Dài",
                                    item["chieu_dai"]
                                )

                                col2.metric(
                                    "Rộng",
                                    item["chieu_rong"]
                                )

                                btn1, btn2 = st.columns(2)

                                btn1.button(
                                    "Sửa",
                                    key=f"m_edit_{index}"
                                )

                                if btn2.button(
                                    "Xóa",
                                    key=f"m_delete_{index}"
                                ):

                                    if os.path.exists(item["images"]):
                                        os.remove(item["images"])

                                    data.pop(index)

                                    with open(
                                        "data.json",
                                        "w",
                                        encoding="utf-8"
                                    ) as f:

                                        json.dump(
                                            data,
                                            f,
                                            ensure_ascii=False,
                                            indent=4
                                        )

                                    st.rerun()
        # ================= DESKTOP =================
        else:

            # st.subheader("Desktop View")

            headers = st.columns([2, 3, 2, 2, 2])

            headers[0].markdown("### Hình")
            headers[1].markdown("### Tên")
            headers[2].markdown("### Dài")
            headers[3].markdown("### Rộng")
            headers[4].markdown("### Action")

            st.divider()

            for index, item in enumerate(data):

                cols = st.columns([2, 3, 2, 2, 2])

                with cols[0]:
                    if os.path.exists(item["images"]):
                        st.image(item["images"], width=160)

                cols[1].write(item["ten"])
                cols[2].write(item["chieu_dai"])
                cols[3].write(item["chieu_rong"])

                with cols[4]:
                    if st.button("Xóa", key=f"delete_{index}"):

                        if os.path.exists(item["images"]):
                            os.remove(item["images"])

                        data.pop(index)

                        with open("data.json", "w", encoding="utf-8") as f:
                            json.dump(data, f, ensure_ascii=False, indent=4)

                        st.rerun()

                st.divider()

    else:
        st.warning("Chưa có dữ liệu")